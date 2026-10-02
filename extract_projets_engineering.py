"""
Réunion engineering : projets Odoo "PRO (LIG)" + "Engineering".

Le script récupère les projets dans Odoo puis ouvre une page web LOCALE
(reunion_ui.html) dans votre navigateur :
- filtres à sélection multiple : étapes, chefs de projet, étiquettes,
  responsables d'actions ; boutons « Tout afficher » / « Effacer les filtres » ;
- par projet : sujets à discuter + actions à entreprendre (case à cocher,
  responsable, échéance) ;
- sauvegarde automatique dans notes_reunion.json (à côté du script) :
  retrouvé à chaque nouvelle extraction ;
- export Excel des projets affichés.

Règles : projets portant LES DEUX étiquettes ; étapes Annulé / Cloturé / Autres /
Template / Canceled exclues (insensible à la casse et aux accents) ; tri par
étape (STAGE_ORDER) -> chef de projet -> numéro de projet.

Identifiants : variables d'environnement ODOO_URL, ODOO_DB, ODOO_USER,
ODOO_PASSWORD, ou fichier .env à côté du script (voir .env.example).
"""
import datetime
import getpass
import io
import json
import os
import re
import secrets
import sys
import threading
import traceback
import unicodedata
import urllib.parse
import webbrowser
import xmlrpc.client
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- Répertoire de travail robuste (fonctionne en double-clic) ---
if getattr(sys, 'frozen', False):
    script_dir = os.path.dirname(sys.executable)
else:
    try:
        script_dir = os.path.dirname(os.path.abspath(__file__))
    except NameError:
        script_dir = os.path.abspath(".")
os.chdir(script_dir)


# ==================================================================
#  PARAMÈTRES
# ==================================================================
def load_dotenv(path=".env"):
    """Charge un fichier .env minimal (KEY=VALUE) et affiche un diagnostic."""
    full = os.path.abspath(path)
    if not os.path.exists(path):
        print(f"ℹ️  Pas de fichier .env trouvé : {full}")
        for alt in (".env.txt", "env", "env.txt", ".env.example.txt"):
            if os.path.exists(alt):
                print(f"⚠️  Mais '{alt}' existe : renommez-le exactement en '.env' "
                      "(activez Affichage > Extensions de noms de fichiers dans l'Explorateur).")
        return
    found = []
    # utf-8-sig : tolère le BOM ajouté par le Bloc-notes
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            if line.lower().startswith("export "):
                line = line[7:]
            key, _, value = line.partition("=")
            key, value = key.strip(), value.strip().strip('"').strip("'")
            if value:
                os.environ[key] = value
                found.append(key)
    print(f"ℹ️  .env lu ({full}) : clés renseignées = {found or 'aucune'}")


load_dotenv()
URL = os.environ.get("ODOO_URL", "https://olsen-engineering.odoo.com")
DB = os.environ.get("ODOO_DB", "mynalios-olsen-main-7388485")
USERNAME = os.environ.get("ODOO_USER", "")
PASSWORD = os.environ.get("ODOO_PASSWORD", "")

EXCLUSIONS = ["Annulé", "Cloturé", "Autres", "Template", "Canceled"]
REQUIRED_TAGS = ["PRO (LIG)", "Engineering"]

# Ordre de tri voulu pour les étapes (stage_id)
STAGE_ORDER = [
    "Facture finale",
    "Réception et CE",
    "Livraison et montage",
    "Atelier",
    "Approvisionnement",
    "Technique / Étude",
    "Kick-off",
    "Nouveau",
]

# Orthographes alternatives d'une même étape (même position de tri).
# "Récepton et CE" (faute présente dans Odoo) passe juste après "Facture finale".
STAGE_SYNONYMS = {
    "Réception et CE": ["Récepton et CE", "Reception et CE"],
}

NOTES_FILE = "notes_reunion.json"
UI_FILE = "reunion_ui.html"
NO_MANAGER = "(non assigné)"

# Numéro de projet attendu en tête du nom : lettre + 2 chiffres + "-" + 5 chiffres
PROJECT_NUMBER_RE = re.compile(r'^([A-Za-z]\d{2}-\d{5})\s*(.*)$')
ISO_DATE_RE = re.compile(r'^\d{4}-\d{2}-\d{2}$')


def normalize(text):
    """Minuscule, sans accents, espaces nettoyés : pour comparer des libellés."""
    text = unicodedata.normalize("NFKD", text or "")
    text = "".join(c for c in text if not unicodedata.combining(c))
    return " ".join(text.lower().split())


STAGE_PRIORITY = {normalize(name): i for i, name in enumerate(STAGE_ORDER)}
for _name, _aliases in STAGE_SYNONYMS.items():
    for _alias in _aliases:
        STAGE_PRIORITY[normalize(_alias)] = STAGE_PRIORITY[normalize(_name)]
EXCLUSIONS_NORMALIZED = {normalize(e) for e in EXCLUSIONS}


def is_excluded_stage(stage_name):
    return normalize(stage_name) in EXCLUSIONS_NORMALIZED


def stage_priority(stage_name):
    """Position de tri d'une étape ; les étapes inconnues vont à la fin."""
    return STAGE_PRIORITY.get(normalize(stage_name), len(STAGE_ORDER))


def split_project_name(display_name):
    """Sépare le nom Odoo en (numéro, description)."""
    if not display_name:
        return "", ""
    match = PROJECT_NUMBER_RE.match(display_name.strip())
    if match:
        return match.group(1), match.group(2).strip(" -:–\t")
    return "", display_name.strip()


def iso_date(value):
    """Valeur Odoo (str 'YYYY-MM-DD[ HH:MM:SS]' ou False) -> 'YYYY-MM-DD' ou ''."""
    if isinstance(value, str) and ISO_DATE_RE.match(value.strip()[:10]):
        return value.strip()[:10]
    return ""


def fr_date(iso):
    """'YYYY-MM-DD' -> 'DD/MM/YYYY' (chaîne vide si invalide)."""
    return f"{iso[8:10]}/{iso[5:7]}/{iso[0:4]}" if ISO_DATE_RE.match(iso or "") else ""


def sort_projects(projects):
    """Étape (ordre imposé ; inconnues à la fin, alphabétiques) -> chef -> numéro."""
    unknown = len(STAGE_ORDER)
    return sorted(projects, key=lambda p: (
        stage_priority(p["stage"]),
        normalize(p["stage"]) if stage_priority(p["stage"]) == unknown else "",
        normalize(p["manager"]),
        p["numero"],
    ))


# ==================================================================
#  ODOO
# ==================================================================
def fetch_projects():
    """Retourne la liste des projets (dicts simples), triés."""
    global USERNAME, PASSWORD
    if not USERNAME or not PASSWORD:
        if not (sys.stdin and sys.stdin.isatty()):
            raise Exception("Identifiants manquants : définissez ODOO_USER et ODOO_PASSWORD "
                            "(variables d'environnement ou fichier .env, voir .env.example).")
        print("ℹ️  Aucun identifiant trouvé (.env / variables d'environnement) : saisie manuelle.")
        USERNAME = USERNAME or input("Identifiant Odoo (email) : ").strip()
        PASSWORD = PASSWORD or getpass.getpass("Mot de passe (invisible à la saisie) : ")

    print("🔌 Connexion à Odoo...")
    common = xmlrpc.client.ServerProxy(f"{URL}/xmlrpc/2/common", allow_none=True)
    uid = common.authenticate(DB, USERNAME, PASSWORD, {})
    if not uid:
        raise Exception("Authentification échouée — vérifiez vos identifiants.")
    models = xmlrpc.client.ServerProxy(f"{URL}/xmlrpc/2/object", allow_none=True)

    def call(model, method, args, kwargs=None):
        return models.execute_kw(DB, uid, PASSWORD, model, method, args, kwargs or {})

    print("✅ Connecté.\n⏳ Résolution des tags requis...")
    tag_records = call('project.tags', 'search_read',
                       [[['name', 'in', REQUIRED_TAGS]]], {'fields': ['id', 'name']})
    tag_ids_by_name = {t['name']: t['id'] for t in tag_records}
    missing = set(REQUIRED_TAGS) - set(tag_ids_by_name)
    if missing:
        raise Exception(f"Tags introuvables dans Odoo : {missing}")

    print("⏳ Récupération des étapes à exclure...")
    stages = call('project.project.stage', 'search_read', [[]], {
        'fields': ['id', 'name'], 'context': {'active_test': False}})
    excluded_stage_ids = [s['id'] for s in stages if is_excluded_stage(s['name'])]

    print("⏳ Récupération des projets...")
    # Une condition par tag => le projet doit porter TOUS les tags requis (ET).
    domain = [['tag_ids', 'in', [tid]] for tid in tag_ids_by_name.values()]
    if excluded_stage_ids:
        domain.append(['stage_id', 'not in', excluded_stage_ids])
    raw_projects = call('project.project', 'search_read', [domain], {
        'fields': ['stage_id', 'display_name', 'partner_id', 'user_id', 'date_start', 'date',
                   'tag_ids']})
    print(f"✅ {len(raw_projects)} projets récupérés.")

    all_tag_ids = sorted({tid for p in raw_projects for tid in p['tag_ids']})
    tag_names = {}
    if all_tag_ids:
        tag_names = {t['id']: t['name'] for t in call(
            'project.tags', 'search_read', [[['id', 'in', all_tag_ids]]], {'fields': ['id', 'name']})}

    projects = []
    for p in raw_projects:
        stage = p['stage_id'][1] if p['stage_id'] else ''
        if is_excluded_stage(stage):  # filet de sécurité côté client
            continue
        numero, description = split_project_name(p["display_name"])
        projects.append({
            "id": p["id"],
            "numero": numero,
            "description": description,
            "client": p["partner_id"][1] if p["partner_id"] else "",
            "stage": stage or "(sans étape)",
            "tags": [tag_names[tid] for tid in p["tag_ids"] if tid in tag_names],
            "manager": p["user_id"][1] if p["user_id"] else "",
            "date_start": iso_date(p["date_start"]),
            "date_end": iso_date(p["date"]),
        })
    return sort_projects(projects)


# ==================================================================
#  NOTES (sujets + actions), sauvegardées dans notes_reunion.json
# ==================================================================
def clean_note(raw):
    """Valide/normalise une note reçue du navigateur. Lève ValueError si invalide."""
    if not isinstance(raw, dict):
        raise ValueError("note invalide")
    sujets = str(raw.get("sujets", ""))[:20000]
    review = str(raw.get("derniere_review", ""))
    review = review if ISO_DATE_RE.match(review) else ""
    actions = []
    for a in (raw.get("actions") or [])[:200]:
        if not isinstance(a, dict):
            continue
        due = str(a.get("due", ""))
        actions.append({
            "id": str(a.get("id", ""))[:40] or secrets.token_hex(4),
            "text": str(a.get("text", ""))[:2000],
            "owner": str(a.get("owner", ""))[:100],
            "done": bool(a.get("done", False)),
            "due": due if ISO_DATE_RE.match(due) else "",
        })
    return {"sujets": sujets, "actions": actions, "derniere_review": review}


def note_has_content(note):
    return bool(note and (note["sujets"].strip() or note["actions"]))


class NotesStore:
    """Projets courants + notes persistées. Thread-safe."""

    def __init__(self, path):
        self.path = path
        self.lock = threading.RLock()
        self.projects = []
        self.data = {"version": 1, "notes": {}, "snapshots": {}, "archived_on": {}}
        self._load()

    def _load(self):
        if not os.path.exists(self.path):
            return
        try:
            with open(self.path, encoding="utf-8") as f:
                loaded = json.load(f)
            for key in self.data:
                if key in loaded:
                    self.data[key] = loaded[key]
        except (OSError, ValueError) as exc:
            # On ne l'écrase surtout pas : on met le fichier de côté.
            backup = self.path + ".illisible"
            os.replace(self.path, backup)
            print(f"⚠️  {self.path} illisible ({exc}) : renommé en {backup}.")

    def _save(self):
        tmp = self.path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.data, f, ensure_ascii=False, indent=1)
        os.replace(tmp, self.path)

    def set_projects(self, projects):
        """Met à jour la liste courante, les instantanés et l'état « archivé »."""
        with self.lock:
            self.projects = projects
            today = datetime.date.today().isoformat()
            current = {str(p["id"]) for p in projects}
            for p in projects:
                self.data["snapshots"][str(p["id"])] = p
                self.data["archived_on"].pop(str(p["id"]), None)
            for pid, note in self.data["notes"].items():
                if pid not in current and note_has_content(note):
                    self.data["archived_on"].setdefault(pid, today)
            self._save()

    def state(self):
        with self.lock:
            current = {str(p["id"]) for p in self.projects}
            return {
                "projects": self.projects,
                "notes": {pid: n for pid, n in self.data["notes"].items() if pid in current},
                "stage_order": STAGE_ORDER,
                "generated_at": datetime.datetime.now().strftime("%d/%m/%Y %H:%M"),
                "archived_count": len(self.archived_rows()),
            }

    def set_note(self, pid, raw):
        note = clean_note(raw)
        with self.lock:
            if pid not in {str(p["id"]) for p in self.projects}:
                raise ValueError("projet inconnu")
            self.data["notes"][pid] = note
            self._save()

    def archived_rows(self):
        with self.lock:
            current = {str(p["id"]) for p in self.projects}
            rows = []
            for pid, since in self.data["archived_on"].items():
                note = self.data["notes"].get(pid)
                snap = self.data["snapshots"].get(pid)
                if pid not in current and snap and note_has_content(note):
                    rows.append((snap, note, since))
            rows.sort(key=lambda r: r[2], reverse=True)
            return rows

    def rows(self, ids=None):
        """[(projet, note)] dans l'ordre d'affichage, filtrés par ids (str) si fourni."""
        with self.lock:
            empty = {"sujets": "", "actions": [], "derniere_review": ""}
            return [(p, self.data["notes"].get(str(p["id"]), empty)) for p in self.projects
                    if ids is None or str(p["id"]) in ids]


# ==================================================================
#  EXPORT EXCEL
# ==================================================================
THIN_BORDER = Border(*([Side(style="thin", color="D9D9D9")] * 4))
BAND_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

EXPORT_HEADERS = ["N° Projet", "Projet", "Client", "Étape", "Étiquettes", "Chef de projet",
                  "Date début", "Date fin", "Sujets à discuter", "Actions à entreprendre",
                  "Dernière review"]
EXPORT_WIDTHS = [12, 30, 22, 20, 24, 18, 12, 12, 40, 55, 14]


def actions_text(actions):
    """☐/☑ texte — responsable (échéance), une action par ligne."""
    lines = []
    for a in actions:
        line = ("☑ " if a["done"] else "☐ ") + (a["text"].strip() or "(sans titre)")
        if a["owner"].strip():
            line += f" — {a['owner'].strip()}"
        if a["due"]:
            line += f" (échéance {fr_date(a['due'])})"
        lines.append(line)
    return "\n".join(lines)


def build_summary(rows):
    """Tableau (étape x chef de projet) : nombre de projets, dans l'ordre des étapes."""
    if not rows:
        return [], []
    managers = sorted({p["manager"] or NO_MANAGER for p, _ in rows}, key=normalize)
    stages = []
    for p, _ in rows:  # déjà trié par étape
        if p["stage"] not in stages:
            stages.append(p["stage"])
    table = []
    for stage in stages:
        counts = [sum(1 for p, _ in rows
                      if p["stage"] == stage and (p["manager"] or NO_MANAGER) == m)
                  for m in managers]
        table.append([stage] + counts + [sum(counts)])
    totals = [sum(row[i] for row in table) for i in range(1, len(managers) + 2)]
    table.append(["Total"] + totals)
    return ["Étape"] + managers + ["Total"], table


def _write_table(ws, headers, widths, lines):
    for c, label in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.row_dimensions[1].height = 30
    for r, line in enumerate(lines, start=2):
        n_lines = 1
        for c, value in enumerate(line, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r % 2 == 0:
                cell.fill = BAND_FILL
            if isinstance(value, str):
                width = widths[c - 1]
                n_lines = max(n_lines, sum(max(1, -(-len(part) // max(width - 2, 1)))
                                           for part in value.split("\n")))
        ws.row_dimensions[r].height = min(max(30, 15 * n_lines), 400)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(lines) + 1, 2)}"


def export_workbook(rows, archived):
    """rows : [(projet, note)] ; archived : [(snapshot, note, date_iso)] -> bytes xlsx."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Réunion"
    _write_table(ws, EXPORT_HEADERS, EXPORT_WIDTHS, [
        [p["numero"], p["description"], p["client"], p["stage"], ", ".join(p["tags"]),
         p["manager"], fr_date(p["date_start"]), fr_date(p["date_end"]),
         n["sujets"], actions_text(n["actions"]), fr_date(n["derniere_review"])]
        for p, n in rows])

    ws_sum = wb.create_sheet("Synthèse")
    header, table = build_summary(rows)
    for c, label in enumerate(header, start=1):
        cell = ws_sum.cell(row=1, column=c, value=label)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for r, line in enumerate(table, start=2):
        for c, value in enumerate(line, start=1):
            cell = ws_sum.cell(row=r, column=c, value=value)
            cell.border = THIN_BORDER
            if c > 1:
                cell.alignment = Alignment(horizontal="center")
            if r == len(table) + 1:
                cell.font = Font(bold=True)
    ws_sum.column_dimensions["A"].width = 24
    for c in range(2, len(header) + 1):
        ws_sum.column_dimensions[get_column_letter(c)].width = 18

    ws_arc = wb.create_sheet("Projets archivés")
    _write_table(ws_arc, EXPORT_HEADERS + ["Archivé le"], EXPORT_WIDTHS + [12], [
        [p["numero"], p["description"], p["client"], p["stage"], ", ".join(p["tags"]),
         p["manager"], fr_date(p["date_start"]), fr_date(p["date_end"]),
         n["sujets"], actions_text(n["actions"]), fr_date(n["derniere_review"]), fr_date(since)]
        for p, n, since in archived])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ==================================================================
#  SERVEUR LOCAL (127.0.0.1 uniquement)
# ==================================================================
def make_handler(store, token, ui_path, refresh_fn, quit_fn):

    class Handler(BaseHTTPRequestHandler):
        server_version = "Reunion/1.0"

        def log_message(self, *args):  # silence
            pass

        # -- utilitaires -------------------------------------------------
        def _host_ok(self):
            host = (self.headers.get("Host") or "").split(":")[0]
            return host in ("127.0.0.1", "localhost")

        def _send(self, code, body, ctype="application/json; charset=utf-8", extra=None):
            if isinstance(body, str):
                body = body.encode("utf-8")
            self.send_response(code)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-store")
            for k, v in (extra or {}).items():
                self.send_header(k, v)
            self.end_headers()
            self.wfile.write(body)

        def _json(self, obj, code=200):
            self._send(code, json.dumps(obj, ensure_ascii=False))

        def _authorized(self, query):
            supplied = self.headers.get("X-Token") or (query.get("t") or [""])[0]
            return secrets.compare_digest(supplied, token)

        # -- routes ------------------------------------------------------
        def do_GET(self):
            if not self._host_ok():
                return self._send(403, "forbidden", "text/plain")
            url = urllib.parse.urlparse(self.path)
            query = urllib.parse.parse_qs(url.query)
            if url.path == "/":
                with open(ui_path, encoding="utf-8") as f:
                    html = f.read().replace("__TOKEN__", token)
                return self._send(200, html, "text/html; charset=utf-8")
            if url.path == "/favicon.ico":
                return self._send(204, b"", "image/x-icon")
            if not self._authorized(query):
                return self._json({"error": "token invalide"}, 403)
            if url.path == "/api/state":
                return self._json(store.state())
            if url.path == "/api/export.xlsx":
                ids = None
                if "ids" in query:
                    ids = {i for i in query["ids"][0].split(",") if i}
                data = export_workbook(store.rows(ids), store.archived_rows())
                name = f"reunion_engineering_{datetime.date.today().isoformat()}.xlsx"
                return self._send(
                    200, data,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    {"Content-Disposition": f'attachment; filename="{name}"'})
            return self._json({"error": "introuvable"}, 404)

        def do_POST(self):
            if not self._host_ok():
                return self._send(403, "forbidden", "text/plain")
            url = urllib.parse.urlparse(self.path)
            if not self._authorized(urllib.parse.parse_qs(url.query)):
                return self._json({"error": "token invalide"}, 403)
            length = int(self.headers.get("Content-Length") or 0)
            if length > 2_000_000:
                return self._json({"error": "trop gros"}, 413)
            try:
                body = json.loads(self.rfile.read(length) or b"{}")
            except ValueError:
                return self._json({"error": "JSON invalide"}, 400)

            if url.path == "/api/note":
                try:
                    store.set_note(str(body.get("id")), body)
                except ValueError as exc:
                    return self._json({"error": str(exc)}, 400)
                return self._json({"ok": True})
            if url.path == "/api/refresh":
                try:
                    refresh_fn()
                except Exception as exc:  # noqa: BLE001 - renvoyé à l'interface
                    traceback.print_exc()
                    return self._json({"error": str(exc)}, 500)
                return self._json(store.state())
            if url.path == "/api/quit":
                self._json({"ok": True})
                quit_fn()
                return None
            return self._json({"error": "introuvable"}, 404)

    return Handler


def serve(store):
    token = secrets.token_urlsafe(24)
    httpd = None

    def refresh():
        store.set_projects(fetch_projects())

    def quit_server():
        threading.Thread(target=httpd.shutdown, daemon=True).start()

    handler = make_handler(store, token, os.path.join(script_dir, UI_FILE), refresh, quit_server)
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    url = f"http://127.0.0.1:{httpd.server_address[1]}/"
    print(f"\n🌐 Page de réunion : {url}")
    print("   (ouverture dans le navigateur ; Ctrl+C ou bouton « Quitter » pour arrêter)")
    webbrowser.open(url)
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        httpd.server_close()
        print("👋 Serveur arrêté. Vos notes sont dans", os.path.abspath(NOTES_FILE))


def main():
    store = NotesStore(NOTES_FILE)
    projects = fetch_projects()
    unknown = {p["stage"] for p in projects if stage_priority(p["stage"]) == len(STAGE_ORDER)}
    if unknown:
        print(f"⚠️  Étape(s) absente(s) de STAGE_ORDER (placées en fin de liste) : {sorted(unknown)}")
        print("   -> Corrigez STAGE_ORDER en haut du script si l'orthographe diffère.")
    store.set_projects(projects)
    serve(store)


if __name__ == "__main__":
    failed = False
    try:
        main()
    except Exception as e:  # noqa: BLE001
        failed = True
        print(f"\n❌ Erreur : {e}")
        traceback.print_exc()
    finally:
        if failed and sys.stdin and sys.stdin.isatty():
            input("\nAppuyez sur Entrée pour fermer...")
